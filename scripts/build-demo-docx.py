#!/usr/bin/env python3
"""
Build docs/ChainWork_Demo_Walkthrough.docx from docs/DEMO_WALKTHROUGH.md.

The Markdown file is the single source of truth; this only re-renders it as a
branded Word document with the screenshots embedded, for handing to a mentor or a
judging panel who would rather not read a repo.

    pip install python-docx
    python scripts/build-demo-docx.py

If a screenshot referenced by the Markdown hasn't been captured yet, a labelled
placeholder is inserted instead of failing, so the document is always buildable.
Run `npm run demo:capture` first for the real thing.
"""

import os
import re
import sys

try:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    sys.exit("python-docx is missing.  pip install python-docx")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS = os.path.join(ROOT, "docs")
SRC = os.path.join(DOCS, "DEMO_WALKTHROUGH.md")
OUT = os.path.join(DOCS, "ChainWork_Demo_Walkthrough.docx")

# ChainWork palette (docs/CLAUDE.md design tokens)
BRONZE = RGBColor(0xA8, 0x5F, 0x2E)
BRONZE_LIGHT = RGBColor(0xD9, 0xA0, 0x66)
INK = RGBColor(0x1A, 0x15, 0x12)
INK2 = RGBColor(0x57, 0x50, 0x47)
INK3 = RGBColor(0x8E, 0x86, 0x7C)
EMERALD = RGBColor(0x1B, 0x8A, 0x5A)

IMG_RE = re.compile(r"^!\[(?P<alt>[^\]]*)\]\((?P<src>[^)]+)\)\s*$")
HEAD_RE = re.compile(r"^(?P<hashes>#{1,6})\s+(?P<text>.*)$")
UL_RE = re.compile(r"^(\s*)[-*]\s+(.*)$")
OL_RE = re.compile(r"^(\s*)(\d+)\.\s+(.*)$")
INLINE_RE = re.compile(
    r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))"
)


# ---------------------------------------------------------------------------
# low-level helpers
# ---------------------------------------------------------------------------
def shade(cell, hex_colour):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hex_colour)
    cell._tc.get_or_add_tcPr().append(el)


def left_bar(paragraph, hex_colour):
    """A coloured left border — used for the demo-vs-production callouts."""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bar = OxmlElement("w:left")
    bar.set(qn("w:val"), "single")
    bar.set(qn("w:sz"), "18")
    bar.set(qn("w:space"), "10")
    bar.set(qn("w:color"), hex_colour)
    borders.append(bar)
    p_pr.append(borders)


def add_inline(paragraph, text, base_size=10.5, colour=INK2, italic=False):
    """Render **bold**, *italic*, `code` and [links](url) into one paragraph."""
    for chunk in INLINE_RE.split(text):
        if not chunk:
            continue
        run = None
        if chunk.startswith("**") and chunk.endswith("**"):
            run = paragraph.add_run(chunk[2:-2])
            run.bold = True
            run.font.color.rgb = INK
        elif chunk.startswith("`") and chunk.endswith("`"):
            run = paragraph.add_run(chunk[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(base_size - 1)
            run.font.color.rgb = BRONZE
        elif chunk.startswith("[") and "](" in chunk:
            label = chunk[1 : chunk.index("](")]
            run = paragraph.add_run(label)
            run.font.color.rgb = BRONZE
            run.underline = True
        elif chunk.startswith("*") and chunk.endswith("*"):
            run = paragraph.add_run(chunk[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(chunk)
        if run.font.size is None:
            run.font.size = Pt(base_size)
        if run.font.color.rgb is None:
            run.font.color.rgb = colour
        if italic:
            run.italic = True
    return paragraph


def heading(doc, level, text):
    sizes = {1: 22, 2: 16, 3: 12.5, 4: 11}
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20 if level <= 2 else 12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    clean = re.sub(r"[*`]", "", text)
    run = p.add_run(clean)
    run.bold = True
    run.font.size = Pt(sizes.get(level, 11))
    run.font.color.rgb = BRONZE if level <= 2 else INK
    if level == 1:
        run.font.name = "Georgia"
    return p


def rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(10)
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "D8D0C4")
    borders.append(bottom)
    p_pr.append(borders)


def add_image(doc, alt, src):
    path = os.path.normpath(os.path.join(DOCS, src))
    if os.path.exists(path):
        doc.add_picture(path, width=Inches(6.3))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        # thin caption under each screenshot
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(14)
        run = cap.add_run(alt)
        run.italic = True
        run.font.size = Pt(8.5)
        run.font.color.rgb = INK3
        return True

    box = doc.add_paragraph()
    box.alignment = WD_ALIGN_PARAGRAPH.CENTER
    box.paragraph_format.space_before = Pt(10)
    box.paragraph_format.space_after = Pt(14)
    run = box.add_run(f"[ screenshot pending — {os.path.basename(src)} ]\n{alt}")
    run.font.size = Pt(9)
    run.italic = True
    run.font.color.rgb = INK3
    return False


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        shade(cell, "F2ECE2")
        p = cell.paragraphs[0]
        run = p.add_run(re.sub(r"[*`]", "", text))
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = INK
    for row in body:
        cells = table.add_row().cells
        for i, text in enumerate(row[: len(header)]):
            p = cells[i].paragraphs[0]
            add_inline(p, text, base_size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def add_code(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.left_indent = Inches(0.25)
    shade_p = OxmlElement("w:shd")
    shade_p.set(qn("w:val"), "clear")
    shade_p.set(qn("w:fill"), "F4EEE5")
    p._p.get_or_add_pPr().append(shade_p)
    run = p.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = INK2


def add_quote(doc, lines):
    """The demo-vs-production callout blocks."""
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.22)
        p.paragraph_format.space_after = Pt(4 if i < len(lines) - 1 else 12)
        left_bar(p, "D9A066")
        add_inline(p, line, base_size=10)


# ---------------------------------------------------------------------------
# cover page
# ---------------------------------------------------------------------------
def cover(doc):
    for _ in range(4):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("CHAINWORK")
    run.font.name = "Georgia"
    run.font.size = Pt(40)
    run.font.color.rgb = BRONZE
    run.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("End-to-End Walkthrough")
    run.font.size = Pt(17)
    run.font.color.rgb = INK

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(16)
    run = p.add_run(
        "One project, start to finish: a business hires an Android developer\n"
        "for ₹1,20,000, split into five phases, paid through blockchain escrow."
    )
    run.font.size = Pt(11.5)
    run.italic = True
    run.font.color.rgb = INK2

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(40)
    run = p.add_run("TESTNET ONLY · NO REAL FUNDS AT RISK")
    run.font.size = Pt(9)
    run.bold = True
    run.font.color.rgb = EMERALD

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "Every screenshot is a real capture of the running application.\n"
        "Every escrow movement is a real blockchain transaction using a test token."
    )
    run.font.size = Pt(9)
    run.font.color.rgb = INK3

    doc.add_page_break()


def footer(doc):
    para = doc.sections[0].footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run("ChainWork — phase escrow for local, time-boxed work · testnet build")
    run.font.size = Pt(8)
    run.font.color.rgb = INK3


# ---------------------------------------------------------------------------
# main render loop
# ---------------------------------------------------------------------------
def build():
    if not os.path.exists(SRC):
        sys.exit(f"Missing {SRC}")
    lines = open(SRC, encoding="utf-8").read().split("\n")

    doc = Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(1.0)
    section.top_margin = section.bottom_margin = Inches(0.9)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK2
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    cover(doc)
    footer(doc)

    i = 0
    missing = 0
    embedded = 0
    in_contents = False

    while i < len(lines):
        line = lines[i].rstrip()

        # skip the markdown-only table of contents
        if line.strip() == "## Contents":
            in_contents = True
            i += 1
            continue
        if in_contents:
            if line.startswith("---"):
                in_contents = False
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        # horizontal rule
        if re.fullmatch(r"-{3,}", line.strip()):
            rule(doc)
            i += 1
            continue

        # fenced code
        if line.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            add_code(doc, block)
            i += 1
            continue

        # heading
        m = HEAD_RE.match(line)
        if m:
            level = len(m.group("hashes"))
            text = re.sub(r"^\d+\.\s*", "", m.group("text"))
            if level == 1:
                i += 1
                continue  # the title lives on the cover page
            heading(doc, level, text)
            i += 1
            continue

        # image
        m = IMG_RE.match(line)
        if m:
            if add_image(doc, m.group("alt"), m.group("src")):
                embedded += 1
            else:
                missing += 1
            i += 1
            continue

        # table
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                add_table(doc, rows)
            continue

        # blockquote
        if line.startswith(">"):
            block = []
            while i < len(lines) and lines[i].startswith(">"):
                content = lines[i].lstrip(">").strip()
                if content:
                    block.append(content)
                i += 1
            add_quote(doc, block)
            continue

        # lists
        m = UL_RE.match(line) or OL_RE.match(line)
        if m:
            ordered = bool(OL_RE.match(line))
            while i < len(lines):
                mm = OL_RE.match(lines[i]) if ordered else UL_RE.match(lines[i])
                if not mm:
                    break
                text = mm.group(3) if ordered else mm.group(2)
                p = doc.add_paragraph(style="List Number" if ordered else "List Bullet")
                p.paragraph_format.space_after = Pt(3)
                add_inline(p, text)
                i += 1
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        # plain paragraph (join soft-wrapped lines)
        block = []
        while i < len(lines) and lines[i].strip() and not re.match(
            r"^(#|!\[|\||>|```|-{3,}|\s*[-*]\s|\s*\d+\.\s)", lines[i]
        ):
            block.append(lines[i].strip())
            i += 1
        text = " ".join(block)
        italic_only = text.startswith("*") and text.endswith("*") and text.count("*") == 2
        p = doc.add_paragraph()
        add_inline(p, text, italic=italic_only, colour=INK3 if italic_only else INK2)

    doc.save(OUT)
    print(f"Wrote {OUT}")
    print(f"  screenshots embedded: {embedded}")
    if missing:
        print(f"  placeholders (not yet captured): {missing}  →  run `npm run demo:capture`")


if __name__ == "__main__":
    build()
